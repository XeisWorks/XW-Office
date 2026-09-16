"""Produktpalette.xlsx -> Product Hub staging importer (PR05).

Reads the commercial Excel workbook and writes it into staging tables
(``models/product_hub_import.py``) for later review/matching. **Never writes to the
workbook** and never touches a canonical product-hub table directly — committing
staged rows is a separate, explicit step (PR06).

Scope (per docs/product_hub/ excel_mapping / PR05):

- ``Produktpalette`` — primary product/price reference. The sheet is not one flat
  table: it is several "EDITION" blocks, each with its own title row, its own
  "Art.Nr./Titel/Beschreibung/SKG*/brutto/netto/netto" header row, then data rows.
  Category comes from the nearest preceding title row containing "EDITION".
- ``Amazon`` — identifier reference (ASIN/FNSKU/ISBN13), merged onto the matching
  Produktpalette-staged row by SKU. Per explicit product decision, Amazon is **never**
  its own category/product family — it becomes a suggested ``Amazon`` tag plus
  ``product_identifier``-shaped rows (scheme ASIN/FNSKU/ISBN13) for the review step.
- ``Besetzungen`` — a flat controlled-vocabulary seed list, recorded on the batch, not
  staged as products.
- ``Händler`` — explicitly a *controlled* reference only (multi-block, non-tabular
  layout); this PR records that the sheet exists but does not parse its structure.
- ``XeisWorks`` / ``MusikHeroes`` — legacy view/relationship reference only, per the
  approved data model (``excel_mapping.XeisWorks``/``.MusikHeroes``); not imported here,
  used later only as grouping evidence (PR06 curated grouping).

Two "netto" columns exist in ``Produktpalette`` with different, undocumented divisors
(``/1.05`` vs ``/1.1``); per the build plan this importer never guesses which one is
correct — both are staged and a mismatch is surfaced as a price warning for manual
review, not silently resolved.
"""
from __future__ import annotations

import logging
import re
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from xw_office.repositories.product_hub_import import ProductHubImportRepository

logger = logging.getLogger(__name__)

_EDITION_MARKER = "edition"
_NETTO_AMBIGUITY_TOLERANCE = Decimal("0.02")
_HEADER_CELL_NOISE = re.compile(r"[.\-\s]")


@dataclass(frozen=True)
class ProduktpaletteRow:
    """One parsed product row from the ``Produktpalette`` sheet."""

    sku: str
    title: str
    description: str
    skg: str
    category: str
    brutto: Decimal | None
    netto_values: tuple[Decimal | None, ...]
    netto_ambiguous: bool


@dataclass(frozen=True)
class AmazonRow:
    """One parsed identifier row from the ``Amazon`` sheet."""

    sku: str
    asin: str
    fnsku: str
    isbn13: str
    title: str


@dataclass
class ExcelImportReport:
    """Outcome of one :meth:`ExcelWorkbookImporter.run` call."""

    batch_id: uuid.UUID
    products_staged: int = 0
    amazon_only_products: int = 0
    identifiers_staged: int = 0
    besetzungen_seed_count: int = 0
    haendler_sheet_present: bool = False
    price_warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _clean_str(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _normalize_header_cell(value: str) -> str:
    return _HEADER_CELL_NOISE.sub("", value).lower()


def _to_decimal(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def parse_produktpalette_rows(rows: list[tuple[Any, ...]]) -> list[ProduktpaletteRow]:
    """Parse the ``Produktpalette`` sheet's raw rows (``ws.iter_rows(values_only=True)``)."""
    results: list[ProduktpaletteRow] = []
    current_category = ""
    in_data_block = False

    for row in rows:
        first = _clean_str(row[0]) if row else ""
        rest_empty = all(_clean_str(cell) == "" for cell in row[1:]) if len(row) > 1 else True

        if not first:
            continue
        if _EDITION_MARKER in first.lower() and rest_empty:
            current_category = first
            in_data_block = False
            continue
        if _normalize_header_cell(first) == "artnr":
            in_data_block = True
            continue
        if not in_data_block:
            continue
        if not first.upper().startswith("XW-"):
            # Footnote/stray text row (e.g. the "*SKG = ..." legend) — not a product row.
            continue

        title = _clean_str(row[1]) if len(row) > 1 else ""
        description = _clean_str(row[2]) if len(row) > 2 else ""
        skg = _clean_str(row[3]) if len(row) > 3 else ""
        brutto = _to_decimal(row[4]) if len(row) > 4 else None
        netto_values = tuple(_to_decimal(row[i]) for i in range(5, len(row)))
        present_netto = [value for value in netto_values if value is not None]
        ambiguous = (
            len(present_netto) >= 2
            and max(present_netto) - min(present_netto) > _NETTO_AMBIGUITY_TOLERANCE
        )

        results.append(
            ProduktpaletteRow(
                sku=first,
                title=title,
                description=description,
                skg=skg,
                category=current_category,
                brutto=brutto,
                netto_values=netto_values,
                netto_ambiguous=ambiguous,
            )
        )
    return results


def parse_amazon_rows(rows: list[tuple[Any, ...]]) -> list[AmazonRow]:
    """Parse the ``Amazon`` sheet's raw rows. Header: ART.-NR./ASIN/FNSKU/ISBN/TITEL."""
    results: list[AmazonRow] = []
    header_seen = False

    for row in rows:
        first = _clean_str(row[0]) if row else ""
        if not first:
            continue
        if _normalize_header_cell(first) == "artnr":
            header_seen = True
            continue
        if not header_seen:
            continue

        asin = _clean_str(row[1]) if len(row) > 1 else ""
        fnsku = _clean_str(row[2]) if len(row) > 2 else ""
        isbn_raw = _clean_str(row[3]) if len(row) > 3 else ""
        isbn13 = re.sub(r"\D", "", isbn_raw)
        title = _clean_str(row[4]) if len(row) > 4 else ""
        results.append(AmazonRow(sku=first, asin=asin, fnsku=fnsku, isbn13=isbn13, title=title))
    return results


def read_besetzungen(rows: list[tuple[Any, ...]]) -> list[str]:
    """Flat controlled-vocabulary seed list (one instrument/ensemble label per row)."""
    return [_clean_str(row[0]) for row in rows if row and _clean_str(row[0])]


class ExcelWorkbookImporter:
    """Read-only ``Produktpalette.xlsx`` -> staging importer."""

    def __init__(self, *, import_repo: ProductHubImportRepository) -> None:
        self._import_repo = import_repo

    def run(self, workbook_path: str | Path) -> ExcelImportReport:
        path = Path(workbook_path)
        import openpyxl  # local import: only this method needs the xlsx engine

        workbook = openpyxl.load_workbook(path, data_only=True)
        batch = self._import_repo.create_batch(
            source="excel", source_metadata={"file_name": path.name}
        )
        report = ExcelImportReport(batch_id=batch.id)

        staged_by_sku: dict[str, uuid.UUID] = {}

        if "Produktpalette" in workbook.sheetnames:
            raw_rows = list(workbook["Produktpalette"].iter_rows(values_only=True))
            for entry in parse_produktpalette_rows(raw_rows):
                try:
                    staged_id = self._stage_produktpalette_row(batch.id, entry)
                except Exception as exc:  # noqa: BLE001 - one bad row must not abort the batch
                    logger.exception("Excel import: Produktpalette row %s failed", entry.sku)
                    report.errors.append(f"Produktpalette {entry.sku}: {exc}")
                    continue
                staged_by_sku[entry.sku.upper()] = staged_id
                report.products_staged += 1
                if entry.netto_ambiguous:
                    report.price_warnings.append(
                        f"{entry.sku}: netto-Werte weichen voneinander ab "
                        f"{entry.netto_values} — manuell prüfen, nicht automatisch übernommen"
                    )

        if "Amazon" in workbook.sheetnames:
            raw_rows = list(workbook["Amazon"].iter_rows(values_only=True))
            for amazon_entry in parse_amazon_rows(raw_rows):
                try:
                    identifiers_added = self._merge_amazon_row(
                        batch.id, amazon_entry, staged_by_sku, report
                    )
                except Exception as exc:  # noqa: BLE001
                    logger.exception("Excel import: Amazon row %s failed", amazon_entry.sku)
                    report.errors.append(f"Amazon {amazon_entry.sku}: {exc}")
                    continue
                report.identifiers_staged += identifiers_added

        if "Besetzungen" in workbook.sheetnames:
            raw_rows = list(workbook["Besetzungen"].iter_rows(values_only=True))
            besetzungen = read_besetzungen(raw_rows)
            report.besetzungen_seed_count = len(besetzungen)
            self._record_source_metadata(batch.id, {"besetzungen": besetzungen})

        if "Händler" in workbook.sheetnames:
            report.haendler_sheet_present = True
            self._record_source_metadata(
                batch.id,
                {
                    "haendler_sheet_present": True,
                    "haendler_note": "multi-block layout, not parsed in PR05 - controlled reference only",
                },
            )

        self._import_repo.finish_batch(
            batch.id,
            status="completed" if not report.errors else "failed",
            error_summary="; ".join(report.errors[:20]),
        )
        return report

    def _stage_produktpalette_row(self, batch_id: uuid.UUID, entry: ProduktpaletteRow) -> uuid.UUID:
        staging_product = self._import_repo.ingest_staging_product(
            import_batch_id=batch_id,
            source="excel",
            source_key=entry.sku.upper(),
            source_external_id=entry.sku,
            sku=entry.sku,
            name=entry.title,
            raw_payload={
                "sheet": "Produktpalette",
                "sku": entry.sku,
                "title": entry.title,
                "description": entry.description,
                "skg": entry.skg,
                "brutto": str(entry.brutto) if entry.brutto is not None else None,
                "netto_values": [str(v) if v is not None else None for v in entry.netto_values],
            },
            normalized_fields={
                "sheet_origin": "Produktpalette",
                "category": entry.category,
                "brutto": str(entry.brutto) if entry.brutto is not None else None,
                "netto_values": [str(v) if v is not None else None for v in entry.netto_values],
                "netto_ambiguous": entry.netto_ambiguous,
            },
        )
        if entry.category:
            self._import_repo.add_category(
                staging_product.id,
                external_category_id=_slugify(entry.category),
                external_category_name=entry.category,
            )
        return staging_product.id

    def _merge_amazon_row(
        self,
        batch_id: uuid.UUID,
        entry: AmazonRow,
        staged_by_sku: dict[str, uuid.UUID],
        report: ExcelImportReport,
    ) -> int:
        key = entry.sku.upper()
        staging_id = staged_by_sku.get(key)
        if staging_id is None:
            staging_product = self._import_repo.ingest_staging_product(
                import_batch_id=batch_id,
                source="excel",
                source_key=key,
                source_external_id=entry.sku,
                sku=entry.sku,
                name=entry.title,
                raw_payload={"sheet": "Amazon", "sku": entry.sku, "title": entry.title},
                normalized_fields={"sheet_origin": "Amazon"},
            )
            staging_id = staging_product.id
            staged_by_sku[key] = staging_id
            report.products_staged += 1
            report.amazon_only_products += 1

        # Amazon becomes a suggested tag, never its own category/product family.
        existing = self._import_repo.get_staging_product(staging_id)
        existing_tags = existing.normalized_fields.get("suggested_tags") if existing else None
        current_tags: set[str] = set(existing_tags) if isinstance(existing_tags, list) else set()
        current_tags.add("Amazon")
        self._import_repo.merge_normalized_fields(
            staging_id, {"suggested_tags": sorted(current_tags)}
        )

        identifiers_added = 0
        for scheme, value in (("ASIN", entry.asin), ("FNSKU", entry.fnsku), ("ISBN13", entry.isbn13)):
            if not value:
                continue
            self._import_repo.add_identifier(
                staging_id, scheme=scheme, value=value, normalized_value=value
            )
            identifiers_added += 1
        return identifiers_added

    def _record_source_metadata(self, batch_id: uuid.UUID, updates: dict[str, object]) -> None:
        batch = self._import_repo.get_batch(batch_id)
        if batch is None:
            return
        merged = dict(batch.source_metadata)
        merged.update(updates)
        self._import_repo.set_batch_metadata(batch_id, merged)


def _slugify(value: str) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text or "category"
