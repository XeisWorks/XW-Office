"""Tests for the Produktpalette.xlsx -> Product Hub staging importer (PR05)."""
from __future__ import annotations

from pathlib import Path

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from xw_office.models.base import Base
from xw_office.repositories.product_hub_import import ProductHubImportRepository
from xw_office.services.product_hub.excel_import import (
    ExcelWorkbookImporter,
    parse_amazon_rows,
    parse_produktpalette_rows,
    read_besetzungen,
)

_REAL_WORKBOOK = Path(__file__).resolve().parents[2] / "docs" / "Produktpalette.xlsx"


# ---------------------------------------------------------------------------
# Pure parsing functions — fixture rows mirror the real sheet layout
# ---------------------------------------------------------------------------


def test_parse_produktpalette_detects_category_from_edition_title() -> None:
    rows = [
        ("TANZLMUSI EDITION",),
        ("Noten für variable Besetzungen",),
        (),
        ("Art.Nr.", "Titel", "Beschreibung", "SKG*", "brutto", "netto", "netto"),
        ("XW-101", "#1 Volksmusik", "Beschreibung 1", 4, 42.9, 40.86, 39.0),
        ("XW-102", "#2 Volksmusik", "Beschreibung 2", 4, 42.9, 40.86, 39.0),
    ]
    parsed = parse_produktpalette_rows(rows)

    assert [row.sku for row in parsed] == ["XW-101", "XW-102"]
    assert all(row.category == "TANZLMUSI EDITION" for row in parsed)
    assert parsed[0].title == "#1 Volksmusik"
    assert parsed[0].skg == "4"


def test_parse_produktpalette_handles_multiple_blocks() -> None:
    rows = [
        ("TANZLMUSI EDITION",),
        (),
        ("Art.Nr.", "Titel"),
        ("XW-101", "Block 1"),
        (),
        ("BLECH4ER EDITION",),
        ("Untertitel",),
        (),
        ("Art.Nr.", "Titel"),
        ("XW-201", "Block 2"),
    ]
    parsed = parse_produktpalette_rows(rows)

    assert [(row.sku, row.category) for row in parsed] == [
        ("XW-101", "TANZLMUSI EDITION"),
        ("XW-201", "BLECH4ER EDITION"),
    ]


def test_parse_produktpalette_skips_footnote_and_non_sku_rows() -> None:
    rows = [
        ("TANZLMUSI EDITION",),
        (),
        ("Art.Nr.", "Titel"),
        ("XW-101", "Gültige Zeile"),
        ("*SKG = Schwierigkeitsgrad", None, "Erläuterung"),
    ]
    parsed = parse_produktpalette_rows(rows)

    assert [row.sku for row in parsed] == ["XW-101"]


def test_parse_produktpalette_rows_outside_any_block_are_ignored() -> None:
    rows = [("XW-999", "Vor jedem Header — darf nicht auftauchen")]
    assert parse_produktpalette_rows(rows) == []


def test_parse_produktpalette_flags_ambiguous_netto_outside_tolerance() -> None:
    rows = [
        ("TANZLMUSI EDITION",),
        (),
        ("Art.Nr.", "Titel", "Beschreibung", "SKG*", "brutto", "netto", "netto"),
        ("XW-101", "Titel", "Beschr.", 4, 42.9, 40.857142857142854, 38.99999999999999),
    ]
    parsed = parse_produktpalette_rows(rows)
    assert parsed[0].netto_ambiguous is True
    assert len(parsed[0].netto_values) == 2


def test_parse_produktpalette_no_ambiguity_when_netto_values_agree() -> None:
    rows = [
        ("TANZLMUSI EDITION",),
        (),
        ("Art.Nr.", "Titel", "Beschreibung", "SKG*", "brutto", "netto", "netto"),
        ("XW-101", "Titel", "Beschr.", 4, 42.9, 39.0, 39.0001),
    ]
    parsed = parse_produktpalette_rows(rows)
    assert parsed[0].netto_ambiguous is False


def test_parse_produktpalette_single_netto_value_is_never_ambiguous() -> None:
    rows = [
        ("TANZLMUSI EDITION",),
        (),
        ("Art.Nr.", "Titel", "Beschreibung", "SKG*", "brutto", "netto"),
        ("XW-101", "Titel", "Beschr.", 4, 42.9, 39.0),
    ]
    parsed = parse_produktpalette_rows(rows)
    assert parsed[0].netto_ambiguous is False


def test_parse_amazon_rows_extracts_asin_fnsku_isbn() -> None:
    rows = [
        ("AMAZON",),
        (),
        (),
        ("ART.-NR.", "ASIN", "FNSKU", "ISBN", "TITEL"),
        ("XW-501.01", 3903426059, "X0019RQ6OP", 9783903426054, "Christkindl-Hits #1 TRP"),
    ]
    parsed = parse_amazon_rows(rows)

    assert len(parsed) == 1
    row = parsed[0]
    assert row.sku == "XW-501.01"
    assert row.asin == "3903426059"
    assert row.fnsku == "X0019RQ6OP"
    assert row.isbn13 == "9783903426054"
    assert row.title == "Christkindl-Hits #1 TRP"


def test_read_besetzungen_returns_flat_list() -> None:
    rows = [("Kleine Böhmische",), (), ("Trp.",), ("Flh1",)]
    assert read_besetzungen(rows) == ["Kleine Böhmische", "Trp.", "Flh1"]


# ---------------------------------------------------------------------------
# ExcelWorkbookImporter — end-to-end against a small real workbook file
# ---------------------------------------------------------------------------


@pytest.fixture
def session_factory() -> sessionmaker[Session]:
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, autoflush=False, expire_on_commit=False, future=True)


@pytest.fixture
def import_repo(session_factory: sessionmaker[Session]) -> ProductHubImportRepository:
    return ProductHubImportRepository(session_factory)


def _write_fixture_workbook(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    produktpalette = workbook.active
    produktpalette.title = "Produktpalette"
    for row in [
        ["TANZLMUSI EDITION"],
        ["Noten für variable Besetzungen"],
        [],
        ["Art.Nr.", "Titel", "Beschreibung", "SKG*", "brutto", "netto", "netto"],
        ["XW-101", "Ohrwürmer #1", "Beschreibung", 4, 42.9, 40.857142857142854, 38.99999999999999],
        ["XW-102", "Ohrwürmer #2", "Beschreibung", 4, 20.0, 18.18, 18.18],
    ]:
        produktpalette.append(row)

    amazon = workbook.create_sheet("Amazon")
    for row in [
        ["AMAZON"],
        [],
        [],
        ["ART.-NR.", "ASIN", "FNSKU", "ISBN", "TITEL"],
        ["XW-101", 3903426059, "X0019RQ6OP", 9783903426054, "Ohrwürmer #1"],
        ["XW-777", 1112223334, "X00AMZONLY", 9781112223330, "Nur bei Amazon"],
    ]:
        amazon.append(row)

    besetzungen = workbook.create_sheet("Besetzungen")
    for row in [["Kleine Böhmische"], ["Trp."], ["Flh1"]]:
        besetzungen.append(row)

    workbook.create_sheet("Händler")

    workbook.save(path)


@pytest.fixture
def fixture_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "Produktpalette.xlsx"
    _write_fixture_workbook(path)
    return path


def test_run_stages_products_with_categories_and_price_warning(
    fixture_workbook: Path, import_repo: ProductHubImportRepository
) -> None:
    importer = ExcelWorkbookImporter(import_repo=import_repo)

    report = importer.run(fixture_workbook)

    assert report.errors == []
    assert len(report.price_warnings) == 1
    assert "XW-101" in report.price_warnings[0]

    staged = {row.sku: row for row in import_repo.list_staging_products()}
    assert staged["XW-101"].normalized_fields["category"] == "TANZLMUSI EDITION"
    categories = import_repo.list_categories(staged["XW-101"].id)
    assert categories[0].external_category_name == "TANZLMUSI EDITION"


def test_run_merges_amazon_identifiers_and_tag_onto_matching_product(
    fixture_workbook: Path, import_repo: ProductHubImportRepository
) -> None:
    importer = ExcelWorkbookImporter(import_repo=import_repo)

    report = importer.run(fixture_workbook)

    staged = {row.sku: row for row in import_repo.list_staging_products()}
    matched = staged["XW-101"]
    # Produktpalette data must survive the later Amazon merge.
    assert matched.normalized_fields["category"] == "TANZLMUSI EDITION"
    assert matched.normalized_fields["suggested_tags"] == ["Amazon"]

    identifiers = import_repo.list_identifiers(matched.id)
    schemes = {row.scheme: row.value for row in identifiers}
    assert schemes == {"ASIN": "3903426059", "FNSKU": "X0019RQ6OP", "ISBN13": "9783903426054"}
    # 3 identifiers for XW-101 plus 3 more for the Amazon-only XW-777 row.
    assert report.identifiers_staged == 6


def test_run_stages_amazon_only_product_without_inventing_category(
    fixture_workbook: Path, import_repo: ProductHubImportRepository
) -> None:
    importer = ExcelWorkbookImporter(import_repo=import_repo)

    report = importer.run(fixture_workbook)

    assert report.amazon_only_products == 1
    staged = {row.sku: row for row in import_repo.list_staging_products()}
    amazon_only = staged["XW-777"]
    assert amazon_only.normalized_fields["sheet_origin"] == "Amazon"
    assert amazon_only.normalized_fields["suggested_tags"] == ["Amazon"]
    # Amazon never becomes a category/product family — only a suggested tag.
    assert import_repo.list_categories(amazon_only.id) == []


def test_run_records_haendler_presence_without_parsing_it(
    fixture_workbook: Path, import_repo: ProductHubImportRepository
) -> None:
    importer = ExcelWorkbookImporter(import_repo=import_repo)

    report = importer.run(fixture_workbook)

    assert report.haendler_sheet_present is True
    batch = import_repo.get_batch(report.batch_id)
    assert batch is not None
    assert batch.source_metadata["haendler_sheet_present"] is True


def test_run_records_besetzungen_seed_on_batch(
    fixture_workbook: Path, import_repo: ProductHubImportRepository
) -> None:
    importer = ExcelWorkbookImporter(import_repo=import_repo)

    report = importer.run(fixture_workbook)

    assert report.besetzungen_seed_count == 3
    batch = import_repo.get_batch(report.batch_id)
    assert batch is not None
    assert batch.source_metadata["besetzungen"] == ["Kleine Böhmische", "Trp.", "Flh1"]


# ---------------------------------------------------------------------------
# Smoke test against the real business workbook (read-only)
# ---------------------------------------------------------------------------


@pytest.mark.skipif(not _REAL_WORKBOOK.exists(), reason="docs/Produktpalette.xlsx not present")
def test_run_against_real_workbook_is_read_only_and_stages_products(
    import_repo: ProductHubImportRepository,
) -> None:
    importer = ExcelWorkbookImporter(import_repo=import_repo)
    before_bytes = _REAL_WORKBOOK.read_bytes()

    report = importer.run(_REAL_WORKBOOK)

    after_bytes = _REAL_WORKBOOK.read_bytes()
    assert before_bytes == after_bytes, "importer must never modify the source workbook"

    assert report.products_staged > 50
    assert report.haendler_sheet_present is True
    assert report.besetzungen_seed_count > 0

    skus = {row.sku for row in import_repo.list_staging_products()}
    assert "XW-101" in skus
